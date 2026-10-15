
from flask import (
    Flask, redirect, render_template, request, render_template_string,
    send_file, url_for
)
from werkzeug.utils import secure_filename
from pathlib import Path
from urllib.parse import quote_plus
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen
from io import BytesIO
from openpyxl import Workbook
from docx import Document
from dotenv import load_dotenv
from pypdf import PdfReader
import json
import os
import re
import sqlite3
import uuid

from candidates import candidates_bp

app = Flask(__name__)
app.secret_key = "change-this-before-deployment"

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_FOLDER = BASE_DIR / "uploads"
DATABASE = BASE_DIR / "hirevanta.db"
load_dotenv(BASE_DIR / ".env")

ALLOWED_EXTENSIONS = {"pdf", "docx", "txt"}
MAX_FILE_SIZE = 10 * 1024 * 1024

UPLOAD_FOLDER.mkdir(exist_ok=True)
app.config["MAX_CONTENT_LENGTH"] = MAX_FILE_SIZE

app.register_blueprint(candidates_bp)


def get_db():
    connection = sqlite3.connect(DATABASE)
    connection.row_factory = sqlite3.Row
    return connection


def initialize_database():
    with get_db() as connection:
        connection.execute("""
            CREATE TABLE IF NOT EXISTS recruitment_searches (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                job_title TEXT NOT NULL,
                description TEXT NOT NULL,
                technical_skills TEXT,
                experience TEXT,
                other_skills TEXT,
                tools TEXT,
                location TEXT,
                education TEXT,
                employment_type TEXT,
                work_mode TEXT,
                uploaded_filename TEXT,
                open_to_work INTEGER NOT NULL DEFAULT 1,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        """)
        columns = {
            row["name"] for row in connection.execute(
                "PRAGMA table_info(recruitment_searches)"
            )
        }
        if "open_to_work" not in columns:
            connection.execute(
                "ALTER TABLE recruitment_searches "
                "ADD COLUMN open_to_work INTEGER NOT NULL DEFAULT 1"
            )
        connection.execute("""
            CREATE TABLE IF NOT EXISTS candidate_search_results (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                search_id INTEGER NOT NULL,
                source TEXT NOT NULL DEFAULT 'LinkedIn',
                profile_title TEXT NOT NULL,
                profile_url TEXT NOT NULL,
                snippet TEXT,
                candidate_name TEXT NOT NULL DEFAULT '',
                source_title TEXT NOT NULL DEFAULT '',
                UNIQUE(search_id, profile_url)
            )
        """)
        result_columns = (
            ("source", "source TEXT NOT NULL DEFAULT 'LinkedIn'"),
            ("candidate_name", "candidate_name TEXT NOT NULL DEFAULT ''"),
            ("source_title", "source_title TEXT NOT NULL DEFAULT ''"),
        )
        for column_name, column_definition in result_columns:
            existing_columns = {
                row["name"] for row in connection.execute(
                    "PRAGMA table_info(candidate_search_results)"
                )
            }
            if column_name not in existing_columns:
                connection.execute(
                    "ALTER TABLE candidate_search_results "
                    f"ADD COLUMN {column_definition}"
                )
        connection.execute("""
            UPDATE candidate_search_results
            SET source_title = profile_title
            WHERE source_title = ''
        """)


initialize_database()


def extract_jd_text(file_path):
    try:
        extension = file_path.suffix.lower()
        if extension == ".txt":
            text = file_path.read_text(encoding="utf-8", errors="replace")
        elif extension == ".pdf":
            reader = PdfReader(str(file_path))
            page_text = []
            for page_index, page in enumerate(reader.pages):
                if page_index >= 50:
                    break
                page_text.append(page.extract_text() or "")
            text = "\n".join(page_text)
        elif extension == ".docx":
            document = Document(str(file_path))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            text = "\n".join(paragraphs + table_cells)
        else:
            return ""
    except Exception:
        return ""

    return text[:20000].strip()


def extract_jd_keywords(description, limit=18):
    stop_words = {
        "about", "across", "after", "also", "and", "are", "as", "at",
        "be", "because", "but", "by", "can", "candidate", "candidates",
        "company", "description", "experience", "for", "from", "have",
        "in", "include", "including", "is", "it", "job", "least", "like",
        "may", "more", "must", "of", "on", "or", "our", "preferred",
        "qualifications", "required", "requirements", "role", "should",
        "skills", "such", "team", "that", "the", "their", "this", "to",
        "we", "will", "with", "work", "you", "your",
    }
    keywords = []
    seen = set()
    for token in re.findall(r"[A-Za-z0-9][A-Za-z0-9+#.-]*", description):
        normalized = token.strip(".-").lower()
        if len(normalized) < 3 or normalized in stop_words or normalized in seen:
            continue
        seen.add(normalized)
        keywords.append(token)
        if len(keywords) >= limit:
            break
    return " ".join(keywords)


def build_platform_queries(search):
    def terms_from_fields(fields, per_field_limit):
        collected = []
        seen = set()
        for field_value in fields:
            for term in extract_jd_keywords(field_value or "", limit=per_field_limit).split():
                normalized = term.lower()
                if normalized not in seen:
                    collected.append(term)
                    seen.add(normalized)
        return collected

    detailed_terms = terms_from_fields([
        search["job_title"], search["technical_skills"],
        search["other_skills"], search["tools"], search["experience"],
        search["location"], search["education"],
        search["employment_type"], search["work_mode"],
    ], per_field_limit=4)
    seen_terms = {term.lower() for term in detailed_terms}
    description_terms = [
        term for term in extract_jd_keywords(search["description"], limit=8).split()
        if term.lower() not in seen_terms
    ]
    detailed_terms.extend(description_terms)
    if search["open_to_work"]:
        detailed_terms.extend(["Open", "to", "Work"])

    fallback_terms = terms_from_fields([
        search["job_title"], search["technical_skills"],
        search["tools"], search["location"],
    ], per_field_limit=4)

    platform_domains = {
        "LinkedIn": "site:linkedin.com/in/",
        "Naukri": "site:naukri.com",
        "Indeed": "site:indeed.com",
    }
    return {
        platform: (
            f"{domain} {' '.join(detailed_terms)}".strip(),
            f"{domain} {' '.join(fallback_terms)}".strip(),
        )
        for platform, domain in platform_domains.items()
    }


def search_profiles_by_platform(platform, query):
    api_key = os.environ.get("SERPAPI_API_KEY", "").strip()
    if not api_key:
        return [], "Set SERPAPI_API_KEY to retrieve profile results automatically."

    parameters = urlencode({
        "engine": "google",
        "q": query,
        "num": 20,
        "api_key": api_key,
    })
    api_request = Request(
        "https://serpapi.com/search.json?" + parameters,
        headers={"Accept": "application/json"}
    )

    try:
        with urlopen(api_request, timeout=20) as response:
            payload = json.loads(response.read())
    except HTTPError as error:
        try:
            error_payload = json.loads(error.read())
            message = str(error_payload.get("error", "Provider rejected the request."))
        except (json.JSONDecodeError, AttributeError):
            message = "Provider rejected the request."
        message = message.replace(api_key, "[redacted]")[:240]
        return [], f"SerpApi returned HTTP {error.code}: {message}"
    except URLError as error:
        return [], f"Could not connect to SerpApi ({type(error.reason).__name__})."
    except Exception as error:
        return [], f"SerpApi request failed ({type(error).__name__})."

    if payload.get("error"):
        message = str(payload["error"]).replace(api_key, "[redacted]")[:240]
        if "hasn't returned any results for this query" in message.lower():
            return [], ""
        return [], f"SerpApi error: {message}"

    results = []
    seen_urls = set()
    for item in payload.get("organic_results", []):
        title = str(item.get("title", "")).strip()
        raw_url = str(item.get("link", "")).strip()
        parsed_url = urlsplit(raw_url)
        hostname = (parsed_url.hostname or "").lower()
        path = parsed_url.path.lower()
        is_platform_profile = False
        if platform == "LinkedIn":
            is_platform_profile = (
                (hostname == "linkedin.com" or hostname.endswith(".linkedin.com"))
                and path.startswith(("/in/", "/pub/"))
            )
        elif platform == "Naukri":
            is_platform_profile = (
                (hostname == "naukri.com" or hostname.endswith(".naukri.com"))
                and path.startswith(("/mnjuser/profile", "/resume/"))
            )
        elif platform == "Indeed":
            is_platform_profile = (
                (hostname == "indeed.com" or hostname.endswith(".indeed.com"))
                and (
                    path.startswith(("/r/", "/profile/"))
                    or hostname == "profile.indeed.com" and path.startswith("/p/")
                )
            )

        if not title or not is_platform_profile:
            continue

        profile_query = ""
        if platform == "Naukri" and path.startswith("/mnjuser/profile"):
            profile_query = urlencode([
                (key, value)
                for key, value in parse_qsl(parsed_url.query)
                if key.lower() == "id"
            ])
        profile_url = urlunsplit((
            "https", hostname, parsed_url.path, profile_query, ""
        ))
        if profile_url in seen_urls:
            continue
        seen_urls.add(profile_url)
        results.append({
            "source": platform,
            "candidate_name": str(
                item.get("candidate_name") or item.get("name") or ""
            ).strip(),
            "profile_title": title,
            "source_title": title,
            "profile_url": profile_url,
            "snippet": str(item.get("snippet", "")).strip(),
        })

    return results[:20], ""


@app.route("/")
def home():
    return render_template("index.html")


@app.route("/create-search", methods=["POST"])
@app.route("/upload-jd", methods=["POST"])
def create_search():
    job_title = request.form.get("job_title", "").strip()
    description = request.form.get("description", "").strip()
    uploaded_file = request.files.get("jd_file")
    uploaded_filename = ""

    if not job_title:
        return "Job title is required.", 400

    if uploaded_file and uploaded_file.filename:
        filename = secure_filename(uploaded_file.filename)

        if "." not in filename:
            return "Invalid filename.", 400

        extension = filename.rsplit(".", 1)[1].lower()
        if extension not in ALLOWED_EXTENSIONS:
            return "Upload a PDF, DOCX or TXT file.", 400

        uploaded_filename = f"{uuid.uuid4().hex}_{filename}"
        saved_path = UPLOAD_FOLDER / uploaded_filename
        uploaded_file.save(saved_path)
        extracted_text = extract_jd_text(saved_path)
        if extracted_text:
            description = "\n".join(
                part for part in [description, extracted_text] if part
            )

    if not description and not uploaded_filename:
        return "Enter a job description or upload a JD file.", 400

    fields = (
        job_title,
        description,
        request.form.get("technical_skills", "").strip(),
        request.form.get("experience", "").strip(),
        request.form.get("other_skills", "").strip(),
        request.form.get("tools", "").strip(),
        request.form.get("location", "").strip(),
        request.form.get("education", "").strip(),
        request.form.get("employment_type", "").strip(),
        request.form.get("work_mode", "").strip(),
        int(request.form.get("open_to_work") == "on"),
        uploaded_filename,
    )

    with get_db() as connection:
        cursor = connection.execute("""
            INSERT INTO recruitment_searches (
                job_title, description, technical_skills, experience,
                other_skills, tools, location, education,
                employment_type, work_mode, open_to_work, uploaded_filename
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """, fields)
        search_id = cursor.lastrowid

    return redirect(url_for("find_candidates", search_id=search_id))


@app.route("/searches")
def saved_searches():
    with get_db() as connection:
        searches = connection.execute("""
            SELECT s.id, s.job_title, s.location, s.created_at,
                   COUNT(r.id) AS result_count
            FROM recruitment_searches AS s
            LEFT JOIN candidate_search_results AS r ON r.search_id = s.id
            GROUP BY s.id
            ORDER BY s.id DESC
        """).fetchall()

    return render_template_string("""
        <h1>Saved Recruitment Searches</h1>
        <ul>
        {% for row in searches %}
            <li>
                <a href="{{ url_for('find_candidates', search_id=row['id']) }}"><strong>{{ row['job_title'] }}</strong></a>
                - {{ row['location'] or 'Location not specified' }}
                ({{ row['result_count'] }} profile results)
            </li>
        {% else %}
            <li>No searches saved yet.</li>
        {% endfor %}
        </ul>
        <a href="/">Create another search</a><br><br>
        <a href="/find-candidates">Find candidates</a>
    """, searches=searches)


@app.route("/find-candidates", methods=["GET", "POST"])
def find_candidates():
    job_title = ""
    skills = ""
    tools = ""
    location = ""
    experience = ""
    description = ""
    other_skills = ""
    education = ""
    employment_type = ""
    work_mode = ""
    open_to_work = True
    edit_mode = request.args.get("edit") == "1"
    searches = []
    results = []
    api_error = ""
    results_by_platform = {
        "LinkedIn": [],
        "Naukri": [],
        "Indeed": [],
    }

    search_id = request.args.get("search_id", type=int)
    if request.method == "GET" and search_id:
        with get_db() as connection:
            saved_search = connection.execute(
                "SELECT * FROM recruitment_searches WHERE id = ?",
                (search_id,)
            ).fetchone()

        if saved_search:
            job_title = saved_search["job_title"]
            description = saved_search["description"] or ""
            if not description and saved_search["uploaded_filename"]:
                saved_path = UPLOAD_FOLDER / saved_search["uploaded_filename"]
                if saved_path.is_file():
                    description = extract_jd_text(saved_path)
                    if description:
                        with get_db() as connection:
                            connection.execute(
                                "UPDATE recruitment_searches SET description = ? WHERE id = ?",
                                (description, search_id)
                            )
            skills = saved_search["technical_skills"] or ""
            other_skills = saved_search["other_skills"] or ""
            tools = saved_search["tools"] or ""
            location = saved_search["location"] or ""
            experience = saved_search["experience"] or ""
            education = saved_search["education"] or ""
            employment_type = saved_search["employment_type"] or ""
            work_mode = saved_search["work_mode"] or ""
            open_to_work = bool(saved_search["open_to_work"])
            if saved_search["uploaded_filename"] and not description:
                api_error = (
                    "No readable JD text was extracted. Scanned PDFs need OCR; "
                    "enter skills and tools manually."
                )
        else:
            return "Recruitment search not found.", 404

    if request.method == "POST":
        submitted_search_id = request.form.get("search_id", type=int)
        job_title = request.form.get("job_title", "").strip()
        description = request.form.get("description", "").strip()
        skills = request.form.get("skills", "").strip()
        other_skills = request.form.get("other_skills", "").strip()
        tools = request.form.get("tools", "").strip()
        location = request.form.get("location", "").strip()
        experience = request.form.get("experience", "").strip()
        education = request.form.get("education", "").strip()
        employment_type = request.form.get("employment_type", "").strip()
        work_mode = request.form.get("work_mode", "").strip()
        open_to_work = request.form.get("open_to_work") == "on"

        if not job_title:
            return "Please enter a job title.", 400

        values = (
            job_title, description, skills, experience, other_skills,
            tools, location, education, employment_type, work_mode,
            int(open_to_work),
        )
        with get_db() as connection:
            if submitted_search_id:
                existing = connection.execute(
                    "SELECT 1 FROM recruitment_searches WHERE id = ?",
                    (submitted_search_id,)
                ).fetchone()
                if not existing:
                    return "Recruitment search not found.", 404
                connection.execute("""
                    UPDATE recruitment_searches
                    SET job_title = ?, description = ?, technical_skills = ?,
                        experience = ?, other_skills = ?, tools = ?, location = ?,
                        education = ?, employment_type = ?, work_mode = ?,
                        open_to_work = ?
                    WHERE id = ?
                """, values + (submitted_search_id,))
                connection.execute(
                    "DELETE FROM candidate_search_results WHERE search_id = ?",
                    (submitted_search_id,)
                )
                search_id = submitted_search_id
            else:
                cursor = connection.execute("""
                    INSERT INTO recruitment_searches (
                        job_title, description, technical_skills, experience,
                        other_skills, tools, location, education,
                        employment_type, work_mode, open_to_work
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """, values)
                search_id = cursor.lastrowid
        return redirect(url_for(
            "find_candidates", search_id=search_id, refresh="1"
            if submitted_search_id else None
        ))

    if search_id:
        with get_db() as connection:
            saved_search = connection.execute(
                "SELECT * FROM recruitment_searches WHERE id = ?",
                (search_id,)
            ).fetchone()
        if not saved_search:
            return "Recruitment search not found.", 404

        job_title = saved_search["job_title"]
        description = saved_search["description"] or ""
        skills = saved_search["technical_skills"] or ""
        other_skills = saved_search["other_skills"] or ""
        tools = saved_search["tools"] or ""
        location = saved_search["location"] or ""
        experience = saved_search["experience"] or ""
        education = saved_search["education"] or ""
        employment_type = saved_search["employment_type"] or ""
        work_mode = saved_search["work_mode"] or ""
        open_to_work = bool(saved_search["open_to_work"])
        platform_queries = build_platform_queries(saved_search)
        searches = [
            (platform, "https://www.google.com/search?q=" + quote_plus(queries[0]))
            for platform, queries in platform_queries.items()
        ]

        with get_db() as connection:
            results = connection.execute("""
                  SELECT source, candidate_name, profile_title, profile_url,
                      snippet, source_title
                FROM candidate_search_results
                WHERE search_id = ?
                ORDER BY id
            """, (search_id,)).fetchall()

        if request.args.get("refresh") == "1" or not results:
            provider_errors = []
            if not os.environ.get("SERPAPI_API_KEY", "").strip():
                provider_errors.append(
                    "SERPAPI_API_KEY is missing. Add it to .env and restart HireVanta."
                )
            else:
                for platform, queries in platform_queries.items():
                    api_results = []
                    provider_error = ""
                    for query in queries:
                        api_results, provider_error = search_profiles_by_platform(
                            platform, query
                        )
                        if provider_error or api_results:
                            break
                    if provider_error:
                        provider_errors.append(f"{platform}: {provider_error}")
                        continue

                    with get_db() as connection:
                        connection.execute(
                            "DELETE FROM candidate_search_results "
                            "WHERE search_id = ? AND source = ?",
                            (search_id, platform)
                        )
                        connection.executemany("""
                            INSERT INTO candidate_search_results (
                                search_id, source, profile_title, profile_url,
                                snippet, candidate_name, source_title
                            ) VALUES (?, ?, ?, ?, ?, ?, ?)
                        """, [
                            (
                                search_id, result["source"],
                                result["profile_title"], result["profile_url"],
                                result["snippet"], result["candidate_name"],
                                result["source_title"],
                            )
                            for result in api_results
                        ])

            with get_db() as connection:
                results = connection.execute("""
                    SELECT source, candidate_name, profile_title, profile_url,
                           snippet, source_title
                    FROM candidate_search_results
                    WHERE search_id = ?
                    ORDER BY id
                """, (search_id,)).fetchall()
            if provider_errors:
                api_error = " ".join(provider_errors)
            elif not results:
                api_error = (
                    "No public profile results were found. Each platform was "
                    "searched with detailed and broader saved-search terms; "
                    "try relaxing requirements or checking the JD text."
                )

        for result in results:
            if result["source"] in results_by_platform:
                results_by_platform[result["source"]].append(result)
        if not results and not api_error:
            api_error = "No public profile results matched these criteria."

    return render_template_string("""
    <!DOCTYPE html>
    <html lang="en">
    <head>
      <meta charset="UTF-8">
      <meta name="viewport" content="width=device-width, initial-scale=1">
      <title>HireVanta | Find Candidates</title>
      <style>
        body { margin: 0; padding: 30px; font-family: Arial, sans-serif;
               background: #f4f6fb; color: #18213a; }
        main { max-width: 850px; margin: auto; }
        section { background: white; padding: 25px; border: 1px solid #e5e8f0;
                  border-radius: 12px; margin-bottom: 20px; }
        label { display: block; margin-bottom: 14px; font-size: 13px;
                font-weight: bold; }
        fieldset { min-width: 0; margin: 0; padding: 0; border: 0; }
                input:not([type="checkbox"]), textarea { display: block; box-sizing: border-box;
          width: 100%; padding: 12px; margin-top: 7px; border: 1px solid #d8ddeb;
          border-radius: 7px; font: inherit; }
                textarea { min-height: 90px; resize: vertical; }
        button, .link { display: inline-block; background: #5145cd; color: white;
          padding: 12px 17px; border: 0; border-radius: 7px;
          text-decoration: none; cursor: pointer; margin: 6px 6px 6px 0; }
        .note { color: #626d82; line-height: 1.6; font-size: 14px; }
        .check { display: flex; align-items: center; gap: 9px; }
      </style>
    </head>
    <body>
    <main>
      <h1>HireVanta</h1>
      <p>Candidate sourcing workspace</p>

      <section>
          <h2>Find Candidates</h2>
        <p class="note">Find up to 20 publicly indexed profile results per platform.
              SerpApi receives the role keywords below, not the full job description.</p>
                <form method="POST">
                    {% if search_id %}
                        <input type="hidden" name="search_id" value="{{ search_id }}">
                        {% if not edit_mode %}<p class="note">Search criteria are saved. Choose Edit search details to change them.</p>{% endif %}
                    {% endif %}
                    <fieldset {% if search_id and not edit_mode %}disabled{% endif %}>
          <label>Job title
            <input name="job_title" required value="{{ job_title }}"
                   placeholder="e.g. Data Analyst">
          </label>
                    <label>Job description or additional JD details
                        <textarea name="description" placeholder="Paste role requirements or JD keywords">{{ description }}</textarea>
                    </label>
          <label>Technical skills
            <input name="skills" value="{{ skills }}"
                   placeholder="e.g. Excel, SQL, Power BI">
          </label>
                    <label>Additional must-have criteria
                        <input name="other_skills" value="{{ other_skills }}"
                                     placeholder="e.g. healthcare domain, AWS certification, team leadership">
                    </label>
          <label>Tools
            <input name="tools" value="{{ tools }}"
                   placeholder="e.g. Tableau, Jira">
          </label>
          <label>Experience
            <input name="experience" value="{{ experience }}"
                   placeholder="e.g. Fresher, 2 years">
          </label>
          <label>Location
            <input name="location" value="{{ location }}"
                   placeholder="e.g. Bengaluru">
          </label>
                    <label>Education
                        <input name="education" value="{{ education }}"
                                     placeholder="e.g. BSc Computer Science">
                    </label>
                    <label>Employment type
                        <input name="employment_type" value="{{ employment_type }}"
                                     placeholder="e.g. Full-time, contract">
                    </label>
                    <label>Work mode
                        <input name="work_mode" value="{{ work_mode }}"
                                     placeholder="e.g. Remote, hybrid, on-site">
                    </label>
          <label class="check">
            <input type="checkbox" name="open_to_work" {% if open_to_work %}checked{% endif %}>
            Prioritize profiles mentioning Open to Work
          </label>
                    </fieldset>
                    {% if search_id %}
                        {% if edit_mode %}
                            <button type="submit">Save changes and rerun search</button>
                            <a class="link" href="{{ url_for('find_candidates', search_id=search_id) }}">Cancel</a>
                        {% else %}
                            <a class="link" href="{{ url_for('find_candidates', search_id=search_id, edit=1) }}">Edit search details</a>
                        {% endif %}
                    {% else %}
                        <button type="submit">Find candidate profiles</button>
                    {% endif %}
        </form>
      </section>

      {% if searches %}
      <section>
                <h2>Candidate profiles</h2>
                <p class="note">Search results can be incomplete or outdated. “Open to Work” is not reliably indexed; verify each profile directly.</p>
                {% if api_error %}<p class="note">{{ api_error }}</p>{% endif %}
                {% if results %}
                    <p><a class="link" href="{{ url_for('export_search_results', search_id=search_id) }}">Download this search's results (Excel)</a></p>
                {% endif %}
                {% for platform, platform_results in results_by_platform.items() %}
                    <h3>{{ platform }} ({{ platform_results|length }})</h3>
                    {% if platform_results %}
                        <ul>
                            {% for result in platform_results %}
                                <li>
                                    <a href="{{ result['profile_url'] }}" target="_blank" rel="noopener noreferrer">{{ result['candidate_name'] or result['source_title'] or result['profile_title'] }}</a>
                                    {% if result['snippet'] %}<p class="note">{{ result['snippet'] }}</p>{% endif %}
                                </li>
                            {% endfor %}
                        </ul>
                    {% else %}
                        <p class="note">No public profile results found for this platform.</p>
                    {% endif %}
                {% endfor %}
                {% if search_id %}<p><a href="{{ url_for('find_candidates', search_id=search_id, refresh=1) }}">Refresh profile results</a></p>{% endif %}
                <details>
                    <summary>Open external search links manually</summary>
                                        <p class="note">These optional links open third-party search pages; they do not populate the profile list above.</p>
                    {% for name, url in searches %}
                      <a class="link" href="{{ url }}" target="_blank"
                         rel="noopener noreferrer">Search {{ name }}</a>
                    {% endfor %}
                </details>
        <p><a href="/candidates">Add a candidate and manage contact details</a></p>
      </section>
      {% endif %}

      <p><a href="/">Create another search</a> |
         <a href="/candidates">Candidate workspace</a> |
         <a href="/searches">Saved searches</a></p>
    </main>
    </body>
    </html>
    """, job_title=job_title, skills=skills, tools=tools,
       location=location, experience=experience, open_to_work=open_to_work,
         description=description, other_skills=other_skills,
         education=education, employment_type=employment_type,
         work_mode=work_mode, searches=searches, results=results, api_error=api_error,
         results_by_platform=results_by_platform, search_id=search_id,
         edit_mode=edit_mode)


def build_search_results_workbook(rows, search_title):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Candidate Results"
    sheet.append([
        "Candidate Name", "Profile Link", "Platform",
        "Search Title", "Source Title"
    ])
    for row in rows:
        sheet.append([
            row["candidate_name"] or "", row["profile_url"], row["source"],
            search_title, row["source_title"] or row["profile_title"]
        ])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 52
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 32
    sheet.column_dimensions["E"].width = 58

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@app.route("/searches/<int:search_id>/results.xlsx")
def export_search_results(search_id):
    with get_db() as connection:
        search = connection.execute(
            "SELECT job_title FROM recruitment_searches WHERE id = ?",
            (search_id,)
        ).fetchone()
        if not search:
            return "Recruitment search not found.", 404
        rows = connection.execute("""
            SELECT source, candidate_name, profile_title, profile_url,
                   source_title
            FROM candidate_search_results
            WHERE search_id = ?
            ORDER BY id
        """, (search_id,)).fetchall()

    output = build_search_results_workbook(rows, search["job_title"])
    return send_file(
        output,
        as_attachment=True,
        download_name=f"HireVanta_Search_{search_id}_Candidates.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


@app.route("/searches/<int:search_id>/selected-results.xlsx", methods=["POST"])
def export_selected_search_results(search_id):
    selected_urls = set(request.form.getlist("profile_urls"))
    if not selected_urls:
        return redirect(url_for(
            "find_candidates", search_id=search_id, selection_error="empty"
        ))

    with get_db() as connection:
        search = connection.execute(
            "SELECT job_title FROM recruitment_searches WHERE id = ?",
            (search_id,)
        ).fetchone()
        if not search:
            return "Recruitment search not found.", 404
        available_rows = connection.execute("""
            SELECT source, candidate_name, profile_title, profile_url,
                   source_title
            FROM candidate_search_results
            WHERE search_id = ?
            ORDER BY id
        """, (search_id,)).fetchall()

    rows = [row for row in available_rows if row["profile_url"] in selected_urls]
    if not rows:
        return redirect(url_for(
            "find_candidates", search_id=search_id, selection_error="invalid"
        ))

    output = build_search_results_workbook(rows, search["job_title"])
    return send_file(
        output,
        as_attachment=True,
        download_name=f"HireVanta_Search_{search_id}_Selected_Candidates.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


@app.route("/candidates/export")
def export_candidates():
    with get_db() as connection:
        rows = connection.execute("""
            SELECT name, email, phone, skills, experience, location,
                   source, profile_url, open_to_work
            FROM candidates
            ORDER BY id DESC
        """).fetchall()

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Candidates"

    sheet.append([
        "Name", "Email", "Phone", "Skills", "Experience",
        "Location", "Source", "Profile URL", "Open to Work"
    ])

    for row in rows:
        sheet.append([row[key] or "" for key in row.keys()])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(
            len(str(cell.value or "")) for cell in column
        ) + 2
        sheet.column_dimensions[letter].width = min(max(width, 14), 35)

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name="HireVanta_Candidates.xlsx",
        mimetype=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )


@app.errorhandler(413)
def file_too_large(error):
    return "File too large. Maximum upload size is 10 MB.", 413


if __name__ == "__main__":
    app.run(debug=True)
